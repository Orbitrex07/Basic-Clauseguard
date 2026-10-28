"""
Script to generate 3 realistic sample contracts in both .docx and .pdf formats
for testing and demonstration in ClauseGuard:
1. standard_commercial_agreement (Balanced / Fair terms)
2. high_risk_vendor_contract (Severe, one-sided, high-risk clauses)
3. mutual_nda (Low-risk mutual non-disclosure agreement)
"""
import os
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

output_dir = os.path.dirname(os.path.abspath(__file__))

CONTRACTS = [
    {
        "id": "standard_commercial_agreement",
        "title": "STANDARD MASTER SERVICES AGREEMENT",
        "preamble": "This Master Services Agreement (\"Agreement\") is entered into as of January 15, 2026 by and between Enterprise Cloud Inc. (\"Client\") and TechCore Solutions LLC (\"Provider\").",
        "sections": [
            {
                "heading": "1. Scope of Services",
                "text": "Provider shall provide the professional cloud infrastructure management and consulting services described in one or more mutually executed Statements of Work (SOW)."
            },
            {
                "heading": "2. Fees and Payment Terms",
                "text": "Client shall pay all undisputed invoices within thirty (30) days of receipt. Invoices shall be submitted monthly in arrears for services actually rendered. Late payments shall accrue interest at 1.0% per month."
            },
            {
                "heading": "3. Term and Renewal",
                "text": "This Agreement shall commence on the Effective Date and remain in effect for an initial period of one (1) year. Thereafter, the Agreement may be renewed upon mutual written agreement of both parties."
            },
            {
                "heading": "4. Termination",
                "text": "Either party may terminate this Agreement without cause upon sixty (60) days prior written notice to the other party. Either party may terminate immediately if the other party commits a material breach and fails to cure such breach within thirty (30) days of written notice."
            },
            {
                "heading": "5. Confidentiality Obligations",
                "text": "Each party agrees to hold the other party's Confidential Information in strict confidence and not disclose it to any third party for a period of three (3) years following disclosure. Confidential Information shall not include information that is publicly known or independently developed."
            },
            {
                "heading": "6. Limitation of Liability",
                "text": "In no event shall either party's aggregate liability under this Agreement exceed the total fees paid or payable by Client in the twelve (12) months preceding the claim. Neither party shall be liable for indirect, incidental, or consequential damages."
            },
            {
                "heading": "7. Mutual Indemnification",
                "text": "Each party shall defend, indemnify, and hold harmless the other party, its officers, and employees against any third-party claims, damages, or reasonable expenses arising out of a material breach of this Agreement or gross negligence."
            },
            {
                "heading": "8. Governing Law and Dispute Resolution",
                "text": "This Agreement shall be governed by and construed in accordance with the laws of the State of Delaware. Any dispute arising out of this Agreement shall be settled by binding arbitration in Wilmington, Delaware."
            }
        ]
    },
    {
        "id": "high_risk_vendor_contract",
        "title": "SOFTWARE VENDOR SUBSCRIPTION AGREEMENT",
        "preamble": "This Subscription Agreement is by and between Global Tech Vendor Corp (\"Vendor\") and Subscriber (\"Customer\").",
        "sections": [
            {
                "heading": "Section 1. Unilateral Termination",
                "text": "Vendor may terminate this Agreement and suspend all access immediately without cause and without prior notice to Customer. Customer shall have no right to terminate this Agreement prior to the expiration of the full multi-year commitment."
            },
            {
                "heading": "Section 2. Unlimited Customer Liability and Total Vendor Disclaimer",
                "text": "Customer shall have unlimited liability for any and all claims, damages, and operational losses arising out of Customer's use of the software. Vendor disclaims all liability and shall have zero liability to Customer under any legal theory, with all warranties explicitly excluded."
            },
            {
                "heading": "Section 3. Perpetual Confidentiality Obligation",
                "text": "Customer shall maintain Vendor's technical documentation and business data in strict secrecy in perpetuity and indefinitely, without limitation of time, surviving any termination or bankruptcy."
            },
            {
                "heading": "Section 4. Sole Negligence Indemnification",
                "text": "Customer agrees to defend, indemnify, and hold harmless Vendor from and against any and all claims, liabilities, judgments, and legal fees, including claims arising solely from Vendor's own negligence, willful misconduct, or system defects."
            },
            {
                "heading": "Section 5. Auto-Renewal Traps",
                "text": "This Agreement shall automatically renew for additional three-year terms unless Customer provides written notice of cancellation exactly five (5) days prior to the expiration date. Failure to provide notice within this strict window results in binding renewal."
            },
            {
                "heading": "Section 6. Discretionary Invoicing and Price Increases",
                "text": "Payment terms shall be determined at Vendor's sole discretion. Vendor reserves the right to increase subscription fees retroactively without limitation and invoice Customer on demand. Payment is due immediately upon receipt."
            },
            {
                "heading": "Section 7. Foreign Dispute Resolution",
                "text": "Any dispute, legal controversy, or claim shall be submitted exclusively to mandatory binding arbitration held in Zurich, Switzerland, under Swiss arbitration rules, with Customer bearing all travel and arbitration costs regardless of outcome."
            }
        ]
    },
    {
        "id": "mutual_nda",
        "title": "MUTUAL NON-DISCLOSURE AGREEMENT",
        "preamble": "This Mutual Non-Disclosure Agreement (\"NDA\") is made between Alpha Dynamics Inc. and Beta Innovations LLC.",
        "sections": [
            {
                "heading": "1. Purpose and Scope",
                "text": "The parties wish to explore a potential business relationship and may disclose confidential proprietary technical and financial information to one another."
            },
            {
                "heading": "2. Confidential Information Definition",
                "text": "Confidential Information refers to any proprietary information marked as confidential or that reasonably ought to be understood as confidential given the nature of the information."
            },
            {
                "heading": "3. Protection and Standard of Care",
                "text": "Each party shall protect the disclosing party's confidential materials using the same degree of care it uses for its own confidential information, but in no case less than a reasonable degree of care."
            },
            {
                "heading": "4. Term of Confidentiality",
                "text": "The confidentiality obligations set forth herein shall survive for a period of two (2) years following the disclosure date, after which obligations shall expire."
            },
            {
                "heading": "5. Standard Exclusions",
                "text": "Confidential Information shall not include information that: (a) is or becomes publicly known without breach of this Agreement; (b) was already in the receiving party's possession prior to disclosure; or (c) is independently developed without reference to the disclosing party's materials."
            },
            {
                "heading": "6. Return or Destruction",
                "text": "Upon written request from the disclosing party, the receiving party shall promptly return or destroy all tangible copies of the Confidential Information."
            },
            {
                "heading": "7. Termination",
                "text": "Either party may terminate discussions and this Agreement upon thirty (30) days prior written notice to the other party."
            },
            {
                "heading": "8. Governing Law",
                "text": "This Agreement shall be governed by the laws of the State of California, without regard to conflict of laws principles."
            }
        ]
    }
]

def generate_docx(contract_data):
    doc = Document()
    doc.add_heading(contract_data["title"], level=0)
    doc.add_paragraph(contract_data["preamble"])
    doc.add_paragraph()

    for sec in contract_data["sections"]:
        doc.add_heading(sec["heading"], level=1)
        doc.add_paragraph(sec["text"])
        doc.add_paragraph()

    filepath = os.path.join(output_dir, f"{contract_data['id']}.docx")
    doc.save(filepath)
    print(f"Generated DOCX: {filepath}")

def generate_pdf(contract_data):
    filepath = os.path.join(output_dir, f"{contract_data['id']}.pdf")
    doc = SimpleDocTemplate(filepath, pagesize=letter, leftMargin=54, rightMargin=54, topMargin=54, bottomMargin=54)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#0f172a'),
        alignment=1,
        spaceAfter=14
    )
    heading_style = ParagraphStyle(
        'DocHeading',
        parent=styles['Heading2'],
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#1e293b'),
        spaceBefore=10,
        spaceAfter=4
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#334155'),
        spaceAfter=8
    )

    story = [
        Paragraph(contract_data["title"], title_style),
        Paragraph(contract_data["preamble"], body_style),
        Spacer(1, 10)
    ]

    for sec in contract_data["sections"]:
        story.append(Paragraph(sec["heading"], heading_style))
        story.append(Paragraph(sec["text"], body_style))
        story.append(Spacer(1, 4))

    doc.build(story)
    print(f"Generated PDF: {filepath}")

def main():
    print("Generating sample contracts...")
    for c in CONTRACTS:
        generate_docx(c)
        generate_pdf(c)
    print("All sample contracts created successfully!")

if __name__ == "__main__":
    main()
