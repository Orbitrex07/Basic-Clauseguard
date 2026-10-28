import os
import pytest
import fitz
import docx
from app.services.parser import DocumentParser, ParsedPage


@pytest.fixture
def sample_pdf(tmp_path):
    pdf_path = str(tmp_path / "test_contract.pdf")
    doc = fitz.open()

    for i in range(3):
        page = doc.new_page()
        # Include repeated header, page body, and repeated footer
        content = (
            "CONFIDENTIAL MASTER SERVICES AGREEMENT\n"
            f"1. Section {i+1}. The Service Provider shall perform consulting services.\n"
            "Payment shall be due Net 30 days.\n"
            f"Page {i+1} of 3\n"
            "Company Internal Document"
        )
        page.insert_text((50, 50), content)

    doc.save(pdf_path)
    doc.close()
    return pdf_path


@pytest.fixture
def sample_docx(tmp_path):
    docx_path = str(tmp_path / "test_contract.docx")
    doc = docx.Document()
    doc.add_heading("Master Services Agreement", level=1)
    doc.add_paragraph("1. Term and Termination: Either party may terminate with 30 days notice.")
    doc.add_paragraph("2. Payment Terms: Invoices are payable Net 30.")
    doc.save(docx_path)
    return docx_path


def test_parse_pdf_and_clean_headers(sample_pdf):
    pages = DocumentParser.parse_pdf(sample_pdf)
    assert len(pages) == 3
    # Check that repeated header "CONFIDENTIAL MASTER SERVICES AGREEMENT" is stripped
    assert "CONFIDENTIAL MASTER SERVICES AGREEMENT" not in pages[0].text
    # Check that substantive content remains
    assert "Section 1" in pages[0].text
    assert "Section 2" in pages[1].text


def test_parse_docx(sample_docx):
    pages = DocumentParser.parse_docx(sample_docx)
    assert len(pages) >= 1
    assert "Term and Termination" in pages[0].text
    assert "Payment Terms" in pages[0].text


def test_whitespace_normalization():
    raw = "Section   1.1\t\t\tScope\n\n\n\n\nAll   services   rendered."
    norm = DocumentParser.normalize_text(raw)
    assert "   " not in norm
    assert "\t" not in norm
    assert "\n\n\n" not in norm
    assert norm.startswith("Section 1.1")
